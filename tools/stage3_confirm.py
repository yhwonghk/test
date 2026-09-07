#!/usr/bin/env python3
"""Stage 3 helper for the CMMP-DI smoke-test pipeline (see AGENTS-STAGE3.md).

    python tools/stage3_confirm.py inventory <TEST_CASE_ID>
    python tools/stage3_confirm.py apply     <TEST_CASE_ID>

inventory  Preflight (artefacts present, B2 == ID, graphify-out/graph.json newer
           than the Stage-1/2 artefacts, ref/sad-index.json staleness), then scan
           EVERY cell of EVERY sheet of the Stage-1/2 workbooks and both
           traceability.md files for "[NEEDS CONFIRMATION - ...]" flags, group them
           by subject, write <ID>-output/stage3/register.json and print the register.

apply      Validate <ID>-output/stage3/resolutions.json against the register, back
           up the artefacts (first run) or restore them from the backup (later
           runs, so re-applying is idempotent), replace the flag text in the
           recorded cells / Markdown table rows, append the Stage 3 section to both
           traceability files, write the confirmation report and run the post-write
           validators. Prints a PASS/FAIL block; exit code 1 on any FAIL.

The script never touches graphify: querying the graph is the agent's job (the
information-source rule in AGENTS-STAGE3.md). The script only reads and writes the
pipeline's own files, and it only ever writes cells that contained a flag.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import shutil
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    print("STOP: openpyxl is not installed (pip install openpyxl)")
    sys.exit(2)

try:  # openpyxl >= 3.1
    from openpyxl.cell.rich_text import CellRichText, TextBlock
except ImportError:  # pragma: no cover
    CellRichText = None  # type: ignore[assignment]
    TextBlock = None  # type: ignore[assignment]

FLAG_RE = re.compile(r"\[NEEDS?\s+CONFIRMATION\s*[-–—:]?\s*(?P<reason>[^\]]*)\]", re.I)
CANDIDATE_RE = re.compile(
    r"candidate\s+(?P<id>[A-Za-z][A-Za-z0-9_./-]*)\s+'(?P<label>[^']*)'(?:\s+score\s+(?P<score>[0-9.]+))?",
    re.I,
)
NOT_FOUND_RE = re.compile(r"not (?:found )?in\b.*?(?:index|graph)|no source found|no sa&d source", re.I)
FEATURE_TAG_RE = re.compile(r"\(\s*Feature\s*:\s*([^)]+?)\s*\)", re.I)
STORY_NO_RE = re.compile(r"Story\s+(\d+)\s+of\s+(\d+)", re.I)
BULLET_RE = re.compile(r"^\s*[•\-\*·]\s*")
TRAILING_QUALIFIER_RE = re.compile(r"\s*\([^()]*\)\s*$")
C6_LINE_RE = re.compile(r"^\s*[•\-\*·]?\s*(?P<name>.+?)\s*\((?P<paren>[^()]*(?:\([^()]*\)[^()]*)*)\)\s*$")

HUMAN_FIELDS = {
    "involved account(s)", "involved accounts", "person in-charge", "person in-charge (steps)",
    "participant", "role", "tester", "testing date", "test case id",
}
HUMAN_WORDS_RE = re.compile(
    r"\b(account|login|post|person in-charge|test case id|role\.xlsx|summary\.xlsx|test\.xlsx|"
    r"participant|department block|no step data)\b",
    re.I,
)
CHANNEL_FIELDS = {"using channel"}
COLUMN_HEADERS = {"using channel", "role", "person in-charge"}  # participant block (rows 5-9)
STATUSES_WRITE = {"RESOLVED", "PARTIAL", "AMBIGUOUS"}
STATUSES_ALL = STATUSES_WRITE | {"OPEN", "HUMAN"}
SELF_REF_RE = re.compile(r"(^|[\\/])(graphify-out|[^\\/]*-output)([\\/]|$)", re.I)
STAGE3_HEADING = "## Stage 3 confirmation pass"


# --------------------------------------------------------------------------- utils
def norm(text: str | None) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def cell_text(value) -> str:
    if value is None:
        return ""
    return str(value)


def today() -> str:
    return _dt.date.today().strftime("%d.%m.%Y")


def short(rel_path: str) -> str:
    """Last two path components, so tests/traceability.md is distinguishable."""
    return "/".join(Path(rel_path).parts[-2:])


def stamp(path: Path) -> str:
    return _dt.datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")


def load_wb(path: Path, rich: bool):
    if rich and CellRichText is not None:
        try:
            return openpyxl.load_workbook(path, rich_text=True)
        except TypeError:  # older openpyxl
            pass
    return openpyxl.load_workbook(path)


class Artefacts:
    def __init__(self, tc_id: str, base: Path):
        self.tc_id = tc_id
        self.base = base
        self.work_dir = base / tc_id
        self.out_dir = base / f"{tc_id}-output"
        self.tests_dir = self.out_dir / "tests"
        self.stage3_dir = self.out_dir / "stage3"
        self.backup_dir = self.out_dir / "stage3-backup"
        self.scenario = self.out_dir / "scenario.xlsx"
        self.trace1 = self.out_dir / "traceability.md"
        self.trace2 = self.tests_dir / "traceability.md"
        self.summary = self.work_dir / "summary.xlsx"
        self.graph = base / "graphify-out" / "graph.json"
        self.index = base / "ref" / "sad-index.json"
        self.stories = sorted(
            p for p in self.tests_dir.glob(f"{tc_id}-story*.xlsx")
            if "Zone.Identifier" not in p.name and not p.name.startswith("~$")
        ) if self.tests_dir.is_dir() else []

    def workbooks(self) -> list[Path]:
        return [self.scenario] + self.stories

    def markdowns(self) -> list[Path]:
        return [p for p in (self.trace1, self.trace2) if p.is_file()]

    def all_files(self) -> list[Path]:
        return [p for p in self.workbooks() + [self.trace1, self.trace2] if p.is_file()]

    def rel(self, path: Path) -> str:
        try:
            return str(path.relative_to(self.base))
        except ValueError:
            return str(path)

    def backup_of(self, path: Path) -> Path:
        return self.backup_dir / path.relative_to(self.out_dir)


# ------------------------------------------------------------------ sheet helpers
def find_label(ws, label: str):
    """First cell whose normalised text starts with the normalised label."""
    target = norm(label)
    for row in ws.iter_rows():
        for c in row:
            t = norm(cell_text(c.value))
            if t and t.startswith(target):
                return c
    return None


def value_cell_right(ws, label_cell):
    """First non-empty cell to the right of a label in the same row (skips merged gaps)."""
    if label_cell is None:
        return None
    for col in range(label_cell.column + 1, ws.max_column + 1):
        c = ws.cell(row=label_cell.row, column=col)
        if cell_text(c.value).strip():
            return c
    return ws.cell(row=label_cell.row, column=label_cell.column + 1)


def header_row_of(ws):
    c = find_label(ws, "Step #")
    return c.row if c is not None else None


def field_for(ws, cell, header_row):
    """Field label of a cell: step-table header, participant-block column header,
    nearest label to the left, or a label a few rows above."""
    if header_row is not None and cell.row > header_row:
        h = ws.cell(row=header_row, column=cell.column)
        text = cell_text(h.value).strip()
        if text:
            return re.sub(r"\s+", " ", text.splitlines()[0]).strip()
    for r in range(cell.row - 1, max(cell.row - 7, 0), -1):
        t = norm(cell_text(ws.cell(row=r, column=cell.column).value))
        if t in COLUMN_HEADERS:
            return re.sub(r"\s+", " ", cell_text(ws.cell(row=r, column=cell.column).value).splitlines()[0]).strip()
    for col in range(cell.column - 1, 0, -1):
        c = ws.cell(row=cell.row, column=col)
        t = cell_text(c.value).strip()
        if t and not FLAG_RE.search(t):
            return re.sub(r"\s+", " ", t.splitlines()[0]).strip()
    # merged label above (e.g. a value block whose label sits one row up)
    for r in range(cell.row - 1, max(cell.row - 3, 0), -1):
        c = ws.cell(row=r, column=cell.column)
        t = cell_text(c.value).strip()
        if t and not FLAG_RE.search(t):
            return re.sub(r"\s+", " ", t.splitlines()[0]).strip()
    return ""


def story_context(ws) -> tuple[str | None, str | None]:
    cover = value_cell_right(ws, find_label(ws, "Scenarios Cover"))
    text = cell_text(cover.value) if cover is not None else ""
    tag = FEATURE_TAG_RE.search(text)
    no = STORY_NO_RE.search(text)
    return (tag.group(1).strip() if tag else None, f"{no.group(1)} of {no.group(2)}" if no else None)


def sheet_id_cell(ws, tc_id: str) -> bool:
    b2 = norm(cell_text(ws["B2"].value))
    if b2 == norm(tc_id):
        return True
    for c in ws[2]:
        if norm(cell_text(c.value)) == norm(tc_id):
            return True
    return False


# ------------------------------------------------------------------- inventory
def classify(field: str, reason: str) -> str:
    f = norm(field)
    if f in HUMAN_FIELDS or f.startswith("involved account") or HUMAN_WORDS_RE.search(reason or ""):
        return "HUMAN"
    if f in CHANNEL_FIELDS:
        return "GRAPH-CONSTRAINED"
    return "GRAPH"


def parse_flag(flag: str) -> dict:
    m = FLAG_RE.search(flag)
    reason = (m.group("reason") if m else "").strip()
    info = {"reason": reason, "kind": "other", "candidate_id": None, "candidate_label": None, "score": None}
    cand = CANDIDATE_RE.search(reason)
    if cand:
        info.update(kind="candidate", candidate_id=cand.group("id"), candidate_label=cand.group("label"),
                    score=float(cand.group("score")) if cand.group("score") else None)
    elif NOT_FOUND_RE.search(reason):
        info["kind"] = "not_found"
    return info


def subject_for(field: str, line: str, flag: str, feature_tag: str | None, info: dict) -> str:
    """What the flag is about. Text in front of the flag wins (a feature bullet);
    a whole-cell generic flag is about its field; a whole-cell specific flag
    ("official system name for 'ME App'") is about its own reason text, so the same
    flag in scenario K5 and in a story's column H lands in one group."""
    before = line[: line.find(flag)] if flag in line else line
    before = BULLET_RE.sub("", before).strip()
    before = re.sub(r"[\(\[]\s*$", "", before).strip()  # drop the opening "(" that wraps the flag
    f = norm(field)
    if f.startswith("covers tender req") and feature_tag:
        return feature_tag
    if before:
        return before
    if info["kind"] == "other" and info["reason"]:
        return info["reason"]
    if feature_tag and f.startswith(("scenario highlight", "testing steps")):
        return f"{field} / {feature_tag}"
    return field or info["reason"]


def scan_workbook(art: Artefacts, path: Path) -> tuple[list[dict], dict]:
    wb = load_wb(path, rich=False)
    hits: list[dict] = []
    meta: dict = {"file": art.rel(path), "id_ok": False, "feature_tag": None, "story": None}
    is_story = path != art.scenario
    for ws in wb.worksheets:
        header_row = header_row_of(ws)
        feature_tag, story_no = story_context(ws) if is_story else (None, None)
        if is_story and feature_tag and not meta["feature_tag"]:
            meta["feature_tag"], meta["story"] = feature_tag, story_no
        if sheet_id_cell(ws, art.tc_id):
            meta["id_ok"] = True
        for row in ws.iter_rows():
            for c in row:
                text = cell_text(c.value)
                if not FLAG_RE.search(text):
                    continue
                field = field_for(ws, c, header_row)
                step_no = cell_text(ws.cell(row=c.row, column=2).value).strip() \
                    if header_row is not None and c.row > header_row else None
                for li, line in enumerate(text.split("\n")):
                    for m in FLAG_RE.finditer(line):
                        flag = m.group(0)
                        info = parse_flag(flag)
                        hits.append({
                            "file": art.rel(path), "sheet": ws.title, "cell": c.coordinate,
                            "line_index": li, "char_offset": m.start(), "flag": flag,
                            "field": field, "context": line.strip()[:300],
                            "step": step_no, "feature_tag": feature_tag, "story": story_no,
                            "subject": subject_for(field, line, flag, feature_tag, info),
                            "class": classify(field, info["reason"]), **info,
                        })
    return hits, meta


def scan_markdown(art: Artefacts, path: Path) -> tuple[list[dict], list[dict]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    hits, open_q = [], []
    in_open, current_story = False, None
    for no, line in enumerate(lines, 1):
        if line.startswith(STAGE3_HEADING):
            break  # a previous Stage 3 section is regenerated by apply, never inventoried
        if line.lstrip().startswith("#"):
            in_open = bool(re.search(r"open question", line, re.I))
            m = re.search(r"\bstory\s+(\d+)\b", line, re.I)
            current_story = m.group(1) if m else None
            continue
        if in_open and line.strip():
            open_q.append({"file": art.rel(path), "line": no, "text": line.strip()})
        is_table = line.lstrip().startswith("|")
        story = current_story
        if is_table:
            first = line.strip().strip("|").split("|")[0].strip()
            if first.isdigit():
                story = first
        for m in FLAG_RE.finditer(line):
            hits.append({"file": art.rel(path), "line": no, "flag": m.group(0), "char_offset": m.start(),
                         "is_table_row": is_table, "story": story, "text": line.strip()[:300],
                         **parse_flag(m.group(0))})
    return hits, open_q


def group_key(subject: str, flag: str) -> tuple[str, str]:
    subj = norm(TRAILING_QUALIFIER_RE.sub("", subject or ""))
    return subj, norm(flag)


def platform_list(art: Artefacts) -> tuple[str | None, list[str], str | None]:
    if not art.summary.is_file():
        return None, [], f"summary.xlsx not found at {art.rel(art.summary)}"
    try:
        wb = load_wb(art.summary, rich=False)
        ws = next((s for s in wb.worksheets if norm(s.title) == "summary"), wb.worksheets[0])
        header = None
        for row in ws.iter_rows():
            if any(norm(cell_text(c.value)) == "test case id" for c in row):
                header = row
                break
        if header is None:
            return None, [], "summary.xlsx: no 'Test Case ID' header found"
        cols = {norm(cell_text(c.value)): c.column for c in header if cell_text(c.value).strip()}
        id_col, plat_col = cols.get("test case id"), cols.get("platform")
        if not id_col or not plat_col:
            return None, [], "summary.xlsx: 'Platform' column not found"
        for r in range(header[0].row + 1, ws.max_row + 1):
            if norm(cell_text(ws.cell(row=r, column=id_col).value)) == norm(art.tc_id):
                raw = cell_text(ws.cell(row=r, column=plat_col).value).strip()
                parts = [p.strip() for p in re.split(r"\s*[+,;\n]\s*", raw) if p.strip()]
                return raw, parts, None
        return None, [], f"summary.xlsx: no row with Test Case ID {art.tc_id}"
    except Exception as exc:  # noqa: BLE001
        return None, [], f"summary.xlsx unreadable: {exc}"


def preflight(art: Artefacts, metas: list[dict]) -> tuple[list[str], list[str], dict]:
    stops, warns, info = [], [], {}
    if not art.scenario.is_file():
        stops.append(f"STOP: Stage-1 output not found: {art.rel(art.scenario)}")
    if not art.trace1.is_file():
        warns.append(f"WARN: {art.rel(art.trace1)} not found - traceability will not be updated")
    if not art.stories:
        warns.append("WARN: no Stage-2 story files under tests/ - resolving the scenario only")
    if art.stories and not art.trace2.is_file():
        warns.append(f"WARN: {art.rel(art.trace2)} not found")
    for m in metas:
        if not m["id_ok"]:
            stops.append(f"STOP: {m['file']} B2 does not equal {art.tc_id}")
        if m["file"] != art.rel(art.scenario) and not m["feature_tag"]:
            warns.append(f"WARN: {m['file']}: no '(Feature: X)' tag found in Scenarios Cover")
    if not art.graph.is_file():
        stops.append(f"STOP: {art.rel(art.graph)} not found - run Stage 0 (/graphify <docs>) first")
    else:
        info["graph_json_mtime"] = stamp(art.graph)
        originals = [art.backup_of(p) if art.backup_of(p).is_file() else p for p in art.all_files()]
        newest = max(originals, key=lambda p: p.stat().st_mtime) if originals else None
        if newest is not None:
            info["newest_artefact"] = art.rel(newest)
            info["newest_artefact_mtime"] = stamp(newest)
            if art.graph.stat().st_mtime < newest.stat().st_mtime:
                stops.append(
                    "STOP: graphify-out/graph.json ({}) is OLDER than {} ({}) - the graph was not updated "
                    "after the Stage-1/2 outputs were written. Run Stage 0 (/graphify <docs> --update) first."
                    .format(info["graph_json_mtime"], info["newest_artefact"], info["newest_artefact_mtime"])
                )
        if art.index.is_file() and art.index.stat().st_mtime < art.graph.stat().st_mtime:
            warns.append(f"WARN: {art.rel(art.index)} ({stamp(art.index)}) is older than graph.json - "
                         "index stale; re-run the Stage 0 index step after Stage 3 (not blocking)")
    if art.backup_dir.is_dir():
        warns.append("WARN: stage3-backup/ exists - this is a re-run; apply will restore the originals first")
    return stops, warns, info


def build_register(art: Artefacts) -> tuple[dict, list[str], list[str]]:
    metas, all_hits = [], []
    for wb_path in art.workbooks():
        if not wb_path.is_file():
            continue
        hits, meta = scan_workbook(art, wb_path)
        metas.append(meta)
        all_hits.extend(hits)
    stops, warns, info = preflight(art, metas)

    groups: dict[tuple[str, str], dict] = {}
    for h in all_hits:
        key = group_key(h["subject"], h["flag"])
        g = groups.get(key)
        if g is None:
            g = groups[key] = {
                "group_id": f"G{len(groups) + 1:02d}", "class": h["class"], "subject": h["subject"],
                "flag": h["flag"], "kind": h["kind"], "reason": h["reason"], "candidate_id": h["candidate_id"],
                "candidate_label": h["candidate_label"], "score": h["score"], "fields": [], "locations": [],
                "md_locations": [], "related_groups": [],
            }
        if h["field"] and h["field"] not in g["fields"]:
            g["fields"].append(h["field"])
        if h["class"] == "HUMAN":
            g["class"] = "HUMAN"
        g["locations"].append({k: h[k] for k in ("file", "sheet", "cell", "line_index", "char_offset",
                                                    "flag", "field", "context", "step", "feature_tag", "story")})
    for k, g in groups.items():
        g["related_groups"] = [o["group_id"] for kk, o in groups.items() if kk != k and kk[0] == k[0]]

    story_tags = {m["story"].split(" ")[0]: norm(TRAILING_QUALIFIER_RE.sub("", m["feature_tag"]))
                  for m in metas if m.get("story") and m.get("feature_tag")}
    unmatched_md, open_questions = [], []
    for md in art.markdowns():
        hits, open_q = scan_markdown(art, md)
        open_questions.extend(open_q)
        for h in hits:
            cands = [g for g in groups.values() if norm(g["flag"]) == norm(h["flag"])]
            if len(cands) > 1:  # generic flag text: disambiguate by subject in the line, then by story
                by_subject = [g for g in cands if norm(TRAILING_QUALIFIER_RE.sub("", g["subject"])) in norm(h["text"])]
                if len(by_subject) == 1:
                    cands = by_subject
                elif h.get("story") in story_tags:
                    tag = story_tags[h["story"]]
                    cands = [g for g in cands if norm(TRAILING_QUALIFIER_RE.sub("", g["subject"])) == tag] or cands
            if len(cands) == 1:
                cands[0]["md_locations"].append(h)
            else:
                h["candidates"] = [g["group_id"] for g in cands]
                unmatched_md.append(h)

    raw_platform, platforms, plat_warn = platform_list(art)
    if plat_warn:
        warns.append(f"WARN: {plat_warn} - Using Channel flags cannot be checked against Platform")
    register = {
        "test_case_id": art.tc_id, "generated_at": _dt.datetime.now().isoformat(timespec="seconds"),
        "files": [m["file"] for m in metas] + [art.rel(p) for p in art.markdowns()],
        "stories": [{"file": m["file"], "story": m["story"], "feature_tag": m["feature_tag"]}
                    for m in metas if m["file"] != art.rel(art.scenario)],
        "platform_raw": raw_platform, "platform_values": platforms, **info,
        "flag_count": len(all_hits), "groups": list(groups.values()),
        "unmatched_md_flags": unmatched_md, "open_questions": open_questions,
    }
    return register, stops, warns


def print_register(reg: dict) -> None:
    print(f"Stage 3 register for {reg['test_case_id']}  ({reg['generated_at']})")
    print(f"  files scanned : {', '.join(reg['files'])}")
    if reg.get("graph_json_mtime"):
        print(f"  graph.json    : {reg['graph_json_mtime']}   newest artefact: "
              f"{reg.get('newest_artefact')} ({reg.get('newest_artefact_mtime')})")
    print(f"  Platform      : {reg['platform_raw']!r} -> {reg['platform_values']}")
    counts: dict[str, int] = {}
    for g in reg["groups"]:
        counts[g["class"]] = counts.get(g["class"], 0) + 1
    print(f"  flags         : {reg['flag_count']} cell flags in {len(reg['groups'])} groups  "
          f"{counts}  | md flags unmatched: {len(reg['unmatched_md_flags'])}  | open questions: "
          f"{len(reg['open_questions'])}")
    print()
    print("| Group | Class | Subject | Kind / candidate | Flag | Locations |")
    print("|---|---|---|---|---|---|")
    for g in reg["groups"]:
        locs = "; ".join(f"{Path(l['file']).name}!{l['cell']}" + (f" (step {l['step']})" if l["step"] else "")
                         for l in g["locations"])
        md_locs = [f"{short(l['file'])}:L{l['line']}" + ("" if l["is_table_row"] else " (note, not edited)")
                   for l in g["md_locations"]]
        if md_locs:
            locs += "; " + "; ".join(md_locs)
        cand = f"{g['kind']}" + (f" {g['candidate_id']} '{g['candidate_label']}' {g['score']}" if g["candidate_id"] else "")
        rel = f" (related: {', '.join(g['related_groups'])})" if g["related_groups"] else ""
        print(f"| {g['group_id']} | {g['class']} | {g['subject']}{rel} | {cand} | {g['flag']} | {locs} |")
    for g in reg["groups"]:
        print(f"\n{g['group_id']} context: {g['locations'][0]['context']}")
    if reg["unmatched_md_flags"]:
        print("\nMarkdown flags not attached to a group (handled only in the appended section):")
        for h in reg["unmatched_md_flags"]:
            print(f"  {h['file']}:L{h['line']}  {h['text']}")
    if reg["open_questions"]:
        print("\nOpen Questions found in traceability:")
        for q in reg["open_questions"]:
            print(f"  {short(q['file'])}:L{q['line']}  {q['text']}")


def archive_previous_pass(art: Artefacts) -> str:
    """Move the completed pass (backup + register/resolutions/report) under stage3-archive/<ts>/.
    The current artefacts — already carrying that pass's resolutions — become the new baseline."""
    ts = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = art.out_dir / "stage3-archive" / ts
    dest.mkdir(parents=True)
    if art.backup_dir.is_dir():
        shutil.move(str(art.backup_dir), str(dest / "stage3-backup"))
    for name in ("register.json", "resolutions.json", "confirmation-report.md"):
        p = art.stage3_dir / name
        if p.is_file():
            shutil.move(str(p), str(dest / name))
    return art.rel(dest)


def cmd_inventory(args) -> int:
    art = Artefacts(args.test_case_id, Path(args.base).resolve())
    report = art.stage3_dir / "confirmation-report.md"
    if report.is_file():
        if not args.new_pass:
            print(f"STOP: a completed Stage 3 pass exists ({art.rel(report)}).\n"
                  "      To FIX that pass: edit stage3/resolutions.json and re-run `apply` (it restores the backup first).\n"
                  "      To start a NEW pass against a newer graph: re-run `inventory` with --new-pass "
                  "(the previous pass is archived under stage3-archive/<timestamp>/).")
            return 1
        print(f"Previous pass archived under {archive_previous_pass(art)}")
    register, stops, warns = build_register(art)
    art.stage3_dir.mkdir(parents=True, exist_ok=True)
    (art.stage3_dir / "register.json").write_text(json.dumps(register, indent=2, ensure_ascii=False), encoding="utf-8")
    for w in warns:
        print(w)
    print_register(register)
    print(f"\nregister written: {art.rel(art.stage3_dir / 'register.json')}")
    if stops:
        print()
        for s in stops:
            print(s)
        return 1
    if not register["groups"]:
        print("\nNo NEEDS CONFIRMATION flags found - nothing to resolve in Stage 3.")
    return 0


# ----------------------------------------------------------------------- apply
class Check:
    def __init__(self):
        self.lines: list[tuple[str, str]] = []

    def add(self, level: str, msg: str):
        self.lines.append((level, msg))

    def ok(self, msg): self.add("PASS", msg)
    def fail(self, msg): self.add("FAIL", msg)
    def warn(self, msg): self.add("WARN", msg)

    @property
    def failed(self) -> bool:
        return any(l == "FAIL" for l, _ in self.lines)

    def dump(self, title: str):
        print(f"\n=== {title} ===")
        for level, msg in self.lines:
            print(f"[{level}] {msg}")
        print(f"=== {'FAIL' if self.failed else 'PASS'} ===")


def validate_resolutions(reg: dict, res: dict, chk: Check) -> dict[str, dict]:
    groups = {g["group_id"]: g for g in reg["groups"]}
    entries = res.get("resolutions", [])
    by_id: dict[str, dict] = {}
    if norm(res.get("test_case_id")) != norm(reg["test_case_id"]):
        chk.fail(f"resolutions.json test_case_id {res.get('test_case_id')!r} != register {reg['test_case_id']!r}")
    for e in entries:
        gid = e.get("group_id")
        if gid not in groups:
            chk.fail(f"{gid}: unknown group_id (not in register)")
            continue
        if gid in by_id:
            chk.fail(f"{gid}: more than one resolution entry")
        by_id[gid] = e
    missing = [gid for gid in groups if gid not in by_id]
    if missing:
        chk.fail(f"no resolution for group(s): {', '.join(missing)}")
    else:
        chk.ok(f"every register group has exactly one resolution ({len(groups)})")

    platforms = {norm(p) for p in reg.get("platform_values", [])}
    for gid, e in by_id.items():
        g = groups[gid]
        status = str(e.get("status", "")).upper()
        if status not in STATUSES_ALL:
            chk.fail(f"{gid}: status {e.get('status')!r} not in {sorted(STATUSES_ALL)}")
            continue
        e["status"] = status
        new_value = str(e.get("new_value") or "")
        evidence = e.get("evidence") or []
        if status in STATUSES_WRITE and g["class"] == "HUMAN":
            chk.fail(f"{gid}: class HUMAN ({g['subject']}) cannot be {status} from the graph - leave it to the human")
        if status == "RESOLVED":
            if not new_value.strip() or FLAG_RE.search(new_value) or "NEEDS CONFIRMATION" in new_value.upper():
                chk.fail(f"{gid}: RESOLVED needs a non-empty new_value without a NEEDS CONFIRMATION flag")
            if not evidence:
                chk.fail(f"{gid}: RESOLVED without evidence")
            confs = [str(ev.get("confidence", "")).upper() for ev in evidence]
            for ev in evidence:
                if not (ev.get("node") and ev.get("src")):
                    chk.fail(f"{gid}: evidence entry needs 'node' and 'src' as printed by graphify: {ev}")
                if SELF_REF_RE.search(str(ev.get("src", ""))):
                    chk.fail(f"{gid}: evidence src {ev.get('src')!r} is a self-reference (*-output/ or graphify-out/)")
            if any(c not in ("EXTRACTED", "INFERRED") for c in confs):
                chk.fail(f"{gid}: RESOLVED evidence must be EXTRACTED or INFERRED (got {confs})")
            if "EXTRACTED" not in confs and not str(e.get("excerpt_quote") or "").strip():
                chk.fail(f"{gid}: INFERRED-only evidence needs a non-empty excerpt_quote")
            if g["class"] == "GRAPH-CONSTRAINED" and platforms and norm(new_value) not in platforms:
                chk.fail(f"{gid}: channel {new_value!r} is not in the summary.xlsx Platform list {sorted(platforms)}")
            if g["kind"] == "candidate" and g["candidate_id"] and norm(new_value) != norm(g["candidate_id"]) \
                    and not str(e.get("note") or "").strip():
                chk.warn(f"{gid}: resolved to {new_value!r}, not the index candidate {g['candidate_id']!r} - note is empty")
        elif status in ("PARTIAL", "AMBIGUOUS"):
            if not re.match(r"^\[NEEDS CONFIRMATION - .+\]$", new_value.strip(), re.I):
                chk.fail(f"{gid}: {status} new_value must be a refined '[NEEDS CONFIRMATION - ...]' flag")
            if not evidence:
                chk.fail(f"{gid}: {status} without evidence")
        elif status == "OPEN":
            if not str(e.get("reason") or "").strip():
                chk.warn(f"{gid}: OPEN without a reason")
        elif status == "HUMAN" and g["class"] != "HUMAN" and not str(e.get("reason") or "").strip():
            chk.warn(f"{gid}: reclassified to HUMAN without a reason")
        if g["class"] != "HUMAN" and status != "HUMAN" and not (e.get("queries") or e.get("vocab_tokens") is not None):
            chk.warn(f"{gid}: no 'queries' / 'vocab_tokens' recorded - the graphify protocol requires them")
    return by_id


def backup_or_restore(art: Artefacts, chk: Check) -> str:
    files = art.all_files()
    if not art.backup_dir.is_dir():
        art.backup_dir.mkdir(parents=True)
        for p in files:
            dst = art.backup_of(p)
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(p, dst)
        chk.ok(f"backup created: {art.rel(art.backup_dir)} ({len(files)} files)")
        return "created"
    restored = 0
    for p in files:
        src = art.backup_of(p)
        if src.is_file():
            shutil.copy2(src, p)
            restored += 1
        else:
            src.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(p, src)
    chk.ok(f"backup exists: restored {restored} file(s) from {art.rel(art.backup_dir)} before applying")
    return "restored"


def replace_in_line(line: str, flag: str, offset: int, new_value: str) -> str | None:
    if line[offset: offset + len(flag)] == flag:
        return line[:offset] + new_value + line[offset + len(flag):]
    idx = line.find(flag)
    if idx < 0:
        return None
    return line[:idx] + new_value + line[idx + len(flag):]


def set_cell_value(cell, old_text: str, new_text: str) -> None:
    value = cell.value
    if CellRichText is not None and isinstance(value, CellRichText):
        # keep run formatting when the change sits inside a single run
        runs = list(value)
        pos = 0
        for i, run in enumerate(runs):
            rtext = run.text if TextBlock is not None and isinstance(run, TextBlock) else str(run)
            end = pos + len(rtext)
            common = 0
            while common < min(len(old_text), len(new_text)) and old_text[common] == new_text[common]:
                common += 1
            tail = 0
            while tail < min(len(old_text), len(new_text)) - common and old_text[-1 - tail] == new_text[-1 - tail]:
                tail += 1
            if pos <= common and len(old_text) - tail <= end:
                local_start, local_end = common - pos, len(old_text) - tail - pos
                new_run_text = rtext[:local_start] + new_text[common: len(new_text) - tail] + rtext[local_end:]
                if TextBlock is not None and isinstance(run, TextBlock):
                    runs[i] = TextBlock(run.font, new_run_text)
                else:
                    runs[i] = new_run_text
                cell.value = CellRichText(runs)
                return
            pos = end
    cell.value = new_text


def apply_workbooks(art: Artefacts, reg: dict, by_id: dict, chk: Check) -> None:
    per_file: dict[str, list[tuple[dict, dict]]] = {}
    for g in reg["groups"]:
        e = by_id[g["group_id"]]
        if e["status"] not in STATUSES_WRITE:
            continue
        for loc in g["locations"]:
            per_file.setdefault(loc["file"], []).append((loc, e))
    for rel, items in per_file.items():
        path = art.base / rel
        wb = load_wb(path, rich=True)
        # apply per cell, descending by (line, offset) so earlier offsets stay valid
        by_cell: dict[tuple[str, str], list[tuple[dict, dict]]] = {}
        for loc, e in items:
            by_cell.setdefault((loc["sheet"], loc["cell"]), []).append((loc, e))
        written = 0
        for (sheet, coord), cell_items in by_cell.items():
            ws = wb[sheet]
            cell = ws[coord]
            old_text = cell_text(cell.value)
            lines = old_text.split("\n")
            for loc, e in sorted(cell_items, key=lambda x: (x[0]["line_index"], x[0]["char_offset"]), reverse=True):
                li = loc["line_index"]
                if li >= len(lines):
                    chk.fail(f"{rel}!{coord}: line {li} missing - file changed since inventory; re-run inventory")
                    continue
                new_line = replace_in_line(lines[li], loc["flag"], loc["char_offset"], e["new_value"])
                if new_line is None:
                    chk.fail(f"{rel}!{coord}: flag not found in cell - file changed since inventory; re-run inventory")
                    continue
                lines[li] = new_line
            new_text = "\n".join(lines)
            if new_text != old_text:
                set_cell_value(cell, old_text, new_text)
                written += 1
        wb.save(path)
        chk.ok(f"{rel}: {written} cell(s) written")


def field_subject(g: dict) -> str:
    fields = " / ".join(g["fields"])
    return fields if norm(fields) == norm(g["subject"]) else f"{fields} — {g['subject']}"


def evidence_str(e: dict) -> str:
    parts = []
    for ev in e.get("evidence") or []:
        loc = f"{ev.get('src', '')}:{ev.get('loc', '')}".rstrip(":")
        parts.append(f"graphify:{ev.get('node', '')} [{str(ev.get('confidence', '')).upper()}] ({loc})")
    return "; ".join(parts) or (e.get("reason") or "")


def stage3_section(reg: dict, by_id: dict, scope_files: set[str], title_scope: str) -> str:
    rows, still_open, human, resolved_n = [], [], [], 0
    for g in reg["groups"]:
        files = {l["file"] for l in g["locations"]} | {l["file"] for l in g["md_locations"]}
        if scope_files and not (files & scope_files):
            continue
        e = by_id[g["group_id"]]
        locs = "; ".join(f"{Path(l['file']).name}!{l['cell']}" for l in g["locations"])
        new_value = e.get("new_value") if e["status"] in STATUSES_WRITE else "(unchanged)"
        rows.append(f"| {g['group_id']} | {field_subject(g)} | {g['flag']} | "
                    f"{e['status']} | {new_value} | {evidence_str(e)} | {locs} |")
        if e["status"] == "RESOLVED":
            resolved_n += 1
        elif e["status"] == "HUMAN":
            human.append(f"- {g['group_id']} {g['subject']}: {e.get('reason', '')} ({locs})")
        else:
            still_open.append(f"- {g['group_id']} {g['subject']}: {e['status']} — "
                              f"{e.get('reason') or e.get('note') or evidence_str(e)} ({locs})")
    out = [
        "", f"{STAGE3_HEADING} ({today()}) — {title_scope}",
        f"Graph: graphify-out/graph.json ({reg.get('graph_json_mtime', 'n/a')}). Groups in scope: {len(rows)}, "
        f"resolved: {resolved_n}. Register: {reg['test_case_id']}-output/stage3/register.json; decisions: "
        f"stage3/resolutions.json; full report: stage3/confirmation-report.md.",
        "", "| Group | Field — Subject | Old flag | Status | New value | Evidence (graphify) | Locations |",
        "|---|---|---|---|---|---|---|", *rows,
    ]
    if still_open:
        out += ["", "### Still open — for the QA Lead", *still_open]
    if human:
        out += ["", "### Human-only items (role.xlsx / summary.xlsx authority — not a graph question)", *human]
    return "\n".join(out) + "\n"


def apply_markdown(art: Artefacts, reg: dict, by_id: dict, chk: Check) -> None:
    for md in art.markdowns():
        rel = art.rel(md)
        # apply restored the pre-pass file first, so a section from THIS pass cannot be
        # present twice; sections from earlier (archived) passes are kept as history.
        lines = md.read_text(encoding="utf-8").splitlines()
        edits: dict[int, list[tuple[dict, dict]]] = {}
        for g in reg["groups"]:
            e = by_id[g["group_id"]]
            if e["status"] not in STATUSES_WRITE:
                continue
            for h in g["md_locations"]:
                if h["file"] == rel and h["is_table_row"]:
                    edits.setdefault(h["line"], []).append((h, e))
        changed = 0
        for no, items in edits.items():
            line = lines[no - 1]
            for h, e in sorted(items, key=lambda x: x[0]["char_offset"], reverse=True):
                new_line = replace_in_line(line, h["flag"], h["char_offset"], e["new_value"])
                if new_line is None:
                    chk.fail(f"{rel}:L{no}: flag not found - file changed since inventory")
                    continue
                line = new_line
                if e["status"] == "RESOLVED":
                    stripped = line.rstrip()
                    if stripped.endswith("|"):
                        line = stripped[:-1].rstrip() + f" ; Stage 3: {evidence_str(e)} |"
            if line != lines[no - 1]:
                lines[no - 1] = line
                changed += 1
        if md == art.trace1:
            scope = {art.rel(art.scenario), rel}
            title = "scenario.xlsx"
        else:
            scope = {art.rel(p) for p in art.stories} | {rel}
            title = "tests/*.xlsx"
        text = "\n".join(lines).rstrip("\n") + "\n" + stage3_section(reg, by_id, scope, title)
        md.write_text(text, encoding="utf-8")
        chk.ok(f"{rel}: {changed} table row(s) updated, Stage 3 section appended")


def write_report(art: Artefacts, reg: dict, res: dict, by_id: dict, chk_pre: Check) -> Path:
    path = art.stage3_dir / "confirmation-report.md"
    counts: dict[str, int] = {}
    for e in by_id.values():
        counts[e["status"]] = counts.get(e["status"], 0) + 1
    lines = [
        f"# Stage 3 confirmation report — {reg['test_case_id']}", "",
        f"- Date: {today()}", f"- graph.json: {reg.get('graph_json_mtime', 'n/a')}",
        f"- Newest Stage-1/2 artefact: {reg.get('newest_artefact')} ({reg.get('newest_artefact_mtime')})",
        f"- Platform (summary.xlsx): {reg.get('platform_raw')!r}",
        f"- Flags: {reg['flag_count']} cell flags in {len(reg['groups'])} groups; statuses: {counts}",
        "", "## Decisions", "",
        "| Group | Class | Locations | Field | Subject | Old flag | Status | New value | Evidence (node [confidence] src:loc) | Commands |",
        "|---|---|---|---|---|---|---|---|---|---|",
    ]
    for g in reg["groups"]:
        e = by_id[g["group_id"]]
        locs = "; ".join(f"{Path(l['file']).name}!{l['cell']}" for l in g["locations"])
        seen: list[str] = []
        for c in list(e.get("queries") or []) + [ev.get("command") for ev in (e.get("evidence") or [])]:
            if c and c not in seen:
                seen.append(c)
        cmds = "<br>".join(f"`{c}`" for c in seen) or "-"
        new_value = e.get("new_value") if e["status"] in STATUSES_WRITE else "(unchanged)"
        lines.append(f"| {g['group_id']} | {g['class']} | {locs} | {' / '.join(g['fields'])} | {g['subject']} | "
                     f"{g['flag']} | {e['status']} | {new_value} | {evidence_str(e)} | {cmds} |")
    for g in reg["groups"]:
        e = by_id[g["group_id"]]
        if e.get("vocab_tokens") is not None or e.get("excerpt_quote") or e.get("note"):
            lines += ["", f"### {g['group_id']} — {g['subject']}"]
            if e.get("vocab_tokens") is not None:
                lines.append(f"- Query expanded to (from graph vocab, {len(e['vocab_tokens'])} tokens): {e['vocab_tokens']}")
            if e.get("excerpt_quote"):
                lines.append(f"- Excerpt quote: {e['excerpt_quote']}")
            if e.get("note"):
                lines.append(f"- Note: {e['note']}")
            if e.get("reason"):
                lines.append(f"- Reason: {e['reason']}")

    def section(title: str, items: list[str]):
        lines.extend(["", f"## {title}", ""])
        lines.extend(f"- {i}" for i in items) if items else lines.append("- none")

    section("Still open — for the QA Lead", [
        f"{g['group_id']} {g['subject']}: {by_id[g['group_id']]['status']} — "
        f"{by_id[g['group_id']].get('reason') or by_id[g['group_id']].get('note') or evidence_str(by_id[g['group_id']])}"
        for g in reg["groups"] if by_id[g["group_id"]]["status"] in ("PARTIAL", "AMBIGUOUS", "OPEN")])
    section("Human-only items (role.xlsx / summary.xlsx / test.xlsx authority)", [
        f"{g['group_id']} {g['subject']}: {by_id[g['group_id']].get('reason', '')}"
        for g in reg["groups"] if by_id[g["group_id"]]["status"] == "HUMAN"])
    section("Open Questions carried in traceability (read, not rewritten)",
            [f"{short(q['file'])}:L{q['line']} {q['text']}" for q in reg.get("open_questions", [])])
    section("Markdown flags not attached to a group", [f"{h['file']}:L{h['line']} {h['text']}" for h in reg.get("unmatched_md_flags", [])])
    section("Proposed additions (NOT written — QA Lead decides)", list(res.get("proposed_additions") or []))
    section("Contradictions with unflagged values (NOT changed)", list(res.get("contradictions") or []))
    section("Graph gaps — Stage 0 follow-up", list(res.get("graph_gaps") or []))
    lines += ["", "## Pre-write validation", ""] + [f"- [{l}] {m}" for l, m in chk_pre.lines]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def post_validate(art: Artefacts, reg: dict, by_id: dict, chk: Check) -> None:
    allowed: dict[str, set[tuple[str, str]]] = {}
    for g in reg["groups"]:
        if by_id[g["group_id"]]["status"] in STATUSES_WRITE:
            for l in g["locations"]:
                allowed.setdefault(l["file"], set()).add((l["sheet"], l["cell"]))
    c6_map: dict[str, str] = {}
    k5_set: set[str] = set()
    for path in art.workbooks():
        if not path.is_file():
            continue
        rel = art.rel(path)
        new_wb, old_wb = load_wb(path, rich=False), load_wb(art.backup_of(path), rich=False)
        if new_wb.sheetnames != old_wb.sheetnames:
            chk.fail(f"{rel}: sheet list changed")
            continue
        changed: set[tuple[str, str]] = set()
        for ws_new in new_wb.worksheets:
            ws_old = old_wb[ws_new.title]
            if sorted(map(str, ws_new.merged_cells.ranges)) != sorted(map(str, ws_old.merged_cells.ranges)):
                chk.warn(f"{rel}!{ws_new.title}: merged ranges differ from the backup")
            max_r, max_c = max(ws_new.max_row, ws_old.max_row), max(ws_new.max_column, ws_old.max_column)
            for r in range(1, max_r + 1):
                for c in range(1, max_c + 1):
                    a, b = cell_text(ws_old.cell(row=r, column=c).value), cell_text(ws_new.cell(row=r, column=c).value)
                    if a != b:
                        changed.add((ws_new.title, f"{get_column_letter(c)}{r}"))
            for cell_ref in ("Missing Feature", "Not ready for review"):
                lab = find_label(ws_old, cell_ref)
                if lab is not None:
                    for col in range(lab.column, ws_old.max_column + 1):
                        coord = f"{get_column_letter(col)}{lab.row}"
                        if (ws_new.title, coord) in changed:
                            chk.fail(f"{rel}!{coord}: review flag row '{cell_ref}' was modified")
        extra = changed - allowed.get(rel, set())
        if extra:
            chk.fail(f"{rel}: cells changed outside the flagged set: {sorted(extra)}")
        else:
            chk.ok(f"{rel}: {len(changed)} changed cell(s), all within the flagged set")
        for (sheet, coord) in allowed.get(rel, set()) - changed:
            chk.warn(f"{rel}!{coord}: expected a change but the cell is unchanged")

        ws0 = new_wb.worksheets[0]
        if path == art.scenario:
            c6 = value_cell_right(ws0, find_label(ws0, "Included Features"))
            if c6 is None:
                chk.warn(f"{rel}: 'Included Features' label not found - C6 contract not checked")
            else:
                bad = []
                for line in cell_text(c6.value).split("\n"):
                    if not line.strip():
                        continue
                    m = C6_LINE_RE.match(line)
                    if m:
                        c6_map[norm(TRAILING_QUALIFIER_RE.sub("", m.group("name")))] = m.group("paren").strip()
                    else:
                        bad.append(line.strip())
                if bad:
                    chk.fail(f"{rel}: Included Features line(s) break the '• <Feature> (<ID>)' contract: {bad}")
                else:
                    chk.ok(f"{rel}: Included Features contract intact ({len(c6_map)} bullets)")
            k5 = value_cell_right(ws0, find_label(ws0, "Involved System(s)"))
            if k5 is not None:
                k5_set = {norm(BULLET_RE.sub("", l)) for l in cell_text(k5.value).split("\n") if l.strip()}
    for path in art.stories:
        rel = art.rel(path)
        ws = load_wb(path, rich=False).worksheets[0]
        tag, _ = story_context(ws)
        if not tag:
            chk.warn(f"{rel}: no Feature tag - K <-> C6 not checked")
            continue
        matches = [k for k in c6_map if k.startswith(norm(tag))] if c6_map else []
        if c6_map and len(matches) != 1:
            chk.fail(f"{rel}: Feature tag {tag!r} matches {len(matches)} Included Features bullet(s): {matches}")
            continue
        expected = c6_map[matches[0]] if matches else None
        hdr = header_row_of(ws)
        if hdr is None:
            chk.warn(f"{rel}: 'Step #' header not found - step table not checked")
            continue
        cols = {norm(cell_text(c.value).splitlines()[0]) if cell_text(c.value).strip() else "": c.column for c in ws[hdr]}
        k_col = next((c for k, c in cols.items() if k.startswith("covers tender req")), None)
        h_col = next((c for k, c in cols.items() if k.startswith("involved system")), None)
        r, mism, h_mism = hdr + 1, [], []
        while r <= ws.max_row and cell_text(ws.cell(row=r, column=2).value).strip():
            if k_col:
                kv = cell_text(ws.cell(row=r, column=k_col).value).strip()
                if kv and not FLAG_RE.search(kv) and expected is not None and not FLAG_RE.search(expected):
                    ids = [norm(x) for x in re.split(r"[,;\n]", kv) if x.strip()]
                    exp_ids = [norm(x) for x in re.split(r"[,;\n]", expected) if x.strip()]
                    if any(i not in exp_ids for i in ids):
                        mism.append(f"row {r}: {kv!r} vs C6 {expected!r}")
            if h_col and k5_set:
                for sysname in re.split(r"[;\n]", cell_text(ws.cell(row=r, column=h_col).value)):
                    s = norm(BULLET_RE.sub("", sysname))
                    if s and not FLAG_RE.search(sysname) and s not in k5_set:
                        h_mism.append(f"row {r}: {sysname.strip()!r}")
            r += 1
        if mism:
            chk.warn(f"{rel}: Covers Tender Req not equal to the C6 ID of {tag!r}: {mism[:5]}")
        else:
            chk.ok(f"{rel}: Covers Tender Req consistent with C6 for {tag!r}")
        if h_mism:
            chk.warn(f"{rel}: Involved System value(s) not in scenario K5: {sorted(set(h_mism))[:5]}")
        chan_hdr = find_label(ws, "Using Channel")
        platforms = {norm(p) for p in reg.get("platform_values", [])}
        if chan_hdr is not None and platforms:
            bad = []
            for rr in range(chan_hdr.row + 1, chan_hdr.row + 5):
                v = cell_text(ws.cell(row=rr, column=chan_hdr.column).value).strip()
                if v and not FLAG_RE.search(v) and norm(v) not in platforms:
                    bad.append(v)
            if bad:
                chk.warn(f"{rel}: Using Channel value(s) not in Platform {sorted(platforms)}: {bad}")
    for md in art.markdowns():
        if STAGE3_HEADING in md.read_text(encoding="utf-8"):
            chk.ok(f"{art.rel(md)}: Stage 3 section present")
        else:
            chk.fail(f"{art.rel(md)}: Stage 3 section missing")
    left = 0
    for path in art.workbooks():
        if path.is_file():
            hits, _ = scan_workbook(art, path)
            left += len(hits)
    chk.ok(f"NEEDS CONFIRMATION flags remaining in workbooks: {left} (was {reg['flag_count']})")


def cmd_apply(args) -> int:
    art = Artefacts(args.test_case_id, Path(args.base).resolve())
    reg_path, res_path = art.stage3_dir / "register.json", art.stage3_dir / "resolutions.json"
    if not reg_path.is_file():
        print(f"STOP: {art.rel(reg_path)} not found - run `inventory` first")
        return 1
    if not res_path.is_file():
        print(f"STOP: {art.rel(res_path)} not found - write the resolutions first (see AGENTS-STAGE3.md)")
        return 1
    reg = json.loads(reg_path.read_text(encoding="utf-8"))
    try:
        res = json.loads(res_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        print(f"STOP: resolutions.json is not valid JSON: {exc}")
        return 1
    pre = Check()
    by_id = validate_resolutions(reg, res, pre)
    pre.dump("Stage 3 pre-write validation")
    if pre.failed:
        print("\nSTOP: fix resolutions.json and re-run apply (nothing was written)")
        return 1

    post = Check()
    backup_or_restore(art, post)
    # the restore may have changed mtimes; graph freshness was already checked in inventory
    apply_workbooks(art, reg, by_id, post)
    apply_markdown(art, reg, by_id, post)
    if post.failed:
        post.dump("Stage 3 write")
        print("\nSTOP: write step failed - re-run `inventory`, then `apply`")
        return 1
    report = write_report(art, reg, res, by_id, pre)
    post_validate(art, reg, by_id, post)
    post.ok(f"report written: {art.rel(report)}")
    counts: dict[str, int] = {}
    for e in by_id.values():
        counts[e["status"]] = counts.get(e["status"], 0) + 1
    post.dump(f"Stage 3 apply + validation for {art.tc_id}  statuses={counts}")
    return 1 if post.failed else 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=["inventory", "apply"])
    ap.add_argument("test_case_id")
    ap.add_argument("--base", default=".", help="pipeline root (default: current directory)")
    ap.add_argument("--new-pass", action="store_true",
                    help="inventory only: archive a completed Stage 3 pass and start a new one on the updated graph")
    args = ap.parse_args(argv)
    return cmd_inventory(args) if args.command == "inventory" else cmd_apply(args)


if __name__ == "__main__":
    sys.exit(main())
